# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, ".")
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

p = r"d:\10_Workspace\translation-research-202610\翻译研究方法\第1组\问卷调查法正反例_v20.pptx"
prs = Presentation(p)
W, H = 13.3333, 7.5

ly = prs.slides[0].slide_layout
print("logo:", sum(1 for sh in ly.shapes if sh.shape_type == 13))

f = ImageFont.truetype(r"C:\Windows\Fonts\simsun.ttc", 20, index=0)
def tw(t): return f.getlength(t) / 96.0

issues = []
for si, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        x = Emu(sh.left).inches if sh.left is not None else 0
        y = Emu(sh.top).inches if sh.top is not None else 0
        w = Emu(sh.width).inches if sh.width is not None else 0
        h = Emu(sh.height).inches if sh.height is not None else 0
        if sh.shape_type == 13:
            if x < -0.01 or x+w > W+0.01 or y+h > H+0.01:
                issues.append("S%02d 图片越界" % (si+1))
            continue
        if x < -0.01 or y < -0.01 or x+w > W+0.01 or y+h > H+0.01:
            issues.append("S%02d 越界" % (si+1))
        if sh.has_text_frame and sh.text_frame.text.strip():
            th = 0.0
            for pp in sh.text_frame.paragraphs:
                runs = [r for r in pp.runs if r.font.size and abs(r.font.size.pt-20) < 0.5]
                if not runs: continue
                t = "".join(r.text for r in pp.runs)
                if not t.strip(): continue
                n = max(1, int(tw(t)/max(w-0.02,0.5))+1)
                th += n*20*1.3/72 + 5/72
            if th > h + 0.15:
                issues.append("S%02d 溢出 %r (需%.2f>框%.2f)" % (si+1, sh.text_frame.text[:12], th, h))

print("问题数:", len(issues))
for it in issues: print("  ", it)
print("总页数:", len(prs.slides._sldIdLst))
