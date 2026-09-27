# -*- coding: utf-8 -*-
"""极简白 · 黑字版（第 2 稿）：问卷调查法 正例/反例，11 页 · 10 分钟。
设计：纯白底 + 黑字（Regular 为主）+ 浅灰辅助；一页一个重点；冰蓝浅灰卡片 + 深蓝左竖条；无边框分栏。
构建后删除所有装饰形状（除带 _keep 标记的卡片），并把每页字数打印出来 —— 超过 120 字/页 就再删。
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from pptx import Presentation
from pptx.util import Pt
from deck_lib import (new_prs, add_slide, para, add_tb, rect, W, H, CR,
                      tw, fit_wrap, EA, MONO)
LAT = "Calibri"          # 西文与标签统一用 Calibri（参考版式：英文接近 Calibri / Arial）

OUT = "问卷调查法_极简白_黑字_10min.pptx"
TOTAL = 11

BLACK = "000000"
BODY = "242424"
MUTED = "707070"
FAINT = "9C9C9C"
ICE = "EEF2F8"
NAVY = "24476E"
X0, CW = 0.62, 12.09

prs = new_prs()
_chars = []


def S():
    return add_slide(prs)


def card(s, x, y, w, h, fill=ICE, bar=0.045):
    """低饱和浅灰蓝卡片 + 左侧深蓝竖条（无边框、无阴影）。"""
    c = rect(s, x, y, w, h, fill=fill)
    c._element.set("_keep", "1")
    b = rect(s, x, y, bar, h, fill=NAVY)
    b._element.set("_keep", "1")
    return c


def line(s, x, y, w, col="E8E8E8", pt=0.75):
    l = rect(s, x, y, w, 0.010, fill=col)
    l._element.set("_keep", "1")
    return l


def txt(s, x, y, w, h, items, align="l", anchor="top"):
    """items = [(text, size, color, bold, font, lh, spc), ...] 或 str/list[str]"""
    tb, tf = add_tb(s, x, y, w, h, anchor=anchor)
    if isinstance(items, str):
        items = [items]
    n = 0
    for it in items:
        if isinstance(it, str):
            it = (it, 12.5, BODY, False, EA, 1.34, None)
        t, sz, col, bold, fnt, lh, spc = (list(it) + [None] * 7)[:7]
        fnt = fnt or EA
        lh = lh or 1.34
        lines, sz2, _ = fit_wrap(t, w, 4.0, sz, lh=lh, bold=bold, min_size=sz - 2.5)
        para(tf, "\n".join(lines), size=sz2, color=col, bold=bold, font=fnt,
             align=align, first=(n == 0), lh=lh, spc=spc)
        n += 1
    return tb, tf


def tag(s, text, en=None, y=0.62):
    t = (text + (" ｜ " + en.upper() if en else "")).strip(" ｜")
    txt(s, X0, y, CW, 0.24, [(t, 9.5, FAINT, False, LAT, 1.2, 1.3)])


def title(s, zh, size=32, y=1.00, w=None):
    w = w or CW
    txt(s, X0, y, w, 0.72, [(zh, size, BLACK, False, EA, 1.16, 0.4)])


def foot(s, no, right=None):
    txt(s, CR - 2.0, H - 0.44, 2.0, 0.22,
        [(right or ("%02d / %02d" % (no, TOTAL - 1)), 8.5, FAINT, False, MONO, 1.2, 0.6)], align="r")


# ============================================================ 00 封面
s = S()
tag(s, "《翻译研究方法》第 1 组 · 方法 A 案例部分", "Questionnaire · Positive vs Negative Cases")
txt(s, X0, 1.62, 11.4, 0.95, [("一份问卷，两种命运", 40, BLACK, True, EA, 1.12, 0.6)])
txt(s, X0, 2.72, 11.4, 0.34,
    [("A Questionnaire, Two Fates ｜ Interpreting × Translation", 14, MUTED, False, LAT, 1.3, None)])
line(s, X0, 4.00, 5.2)
txt(s, X0, 4.22, 11.9, 1.0, [
    ("正例（口译）｜ 周金华、董燕萍 2019，《外语教学与研究》51(6)", 12, BODY, False, EA, 1.6, None),
    ("反例（笔译）｜ Yang & Mustafa 2022, Human Behavior and Emerging Technologies", 12, BODY, False, EA, 1.6, None)])
txt(s, X0, 5.52, 11.6, 0.32,
    [("全程只用一把尺：研究问题 → 构念与题项 → 抽样与施测 → 统计与结论", 12, MUTED, False, EA, 1.4, 0.4)])
txt(s, X0, 6.28, 11.9, 0.7, [
    ("小组分工（4 人）", 9.5, FAINT, False, EA, 1.3, 0.8),
    ("于萍［组长］框架与逻辑主线、正反例评析（P04、P07）｜ 陈冠臻 正例文献核对与流程梳理（P02–P03）", 10, MUTED, False, EA, 1.45, None),
    ("熊芮 反例文献核对与数据摘记（P05–P06）｜ 方燕 对照表、自检清单、思考题与资料汇总（P08–P10）｜ 汇报人：＿＿",
     10, MUTED, False, EA, 1.45, None)])
foot(s, 0, "封面 · 不计页码 ｜ Cover")

# ============================================================ 01 一把尺
s = S()
tag(s, "四步", "The Chain")
title(s, "先看链条", y=1.02)
txt(s, X0, 1.86, 11.4, 0.32, [("闭住就是好问卷；断在任何一环，结论就悬空。", 12.5, MUTED, False, EA, 1.4, None)])
chain = [("研究问题", "要回答什么", "Research question"),
         ("构念与题项", "怎么变成可答的题", "Construct & items"),
         ("抽样与施测", "向谁发、怎么收", "Sampling & admin."),
         ("统计与结论", "证据撑到哪一步", "Analysis & claim")]
bw, g = 2.82, 0.27
for i, (a, b, c) in enumerate(chain):
    x = X0 + i * (bw + g)
    card(s, x, 2.58, bw, 1.72)
    txt(s, x + 0.30, 2.86, bw - 0.56, 1.3, [
        ("%02d" % (i + 1), 11, NAVY, False, MONO, 1.3, 0.6),
        (a, 17, BLACK, False, EA, 1.34, 0.3),
        (b, 11.5, MUTED, False, EA, 1.45, None),
        (c, 10, FAINT, False, LAT, 1.4, None)])
txt(s, X0, 4.66, 11.4, 0.32,
    [("断点最常出现在第 2 与第 4 步：题项测的不是那个构念；结论比数据多走了几步。", 11.5, MUTED, False, EA, 1.4, None)])
foot(s, 1)

# ============================================================ 02 正例概况
s = S()
tag(s, "正例 · 口译研究", "Positive Case")
title(s, "把“熟练度”变成可测的构念", y=1.02)
txt(s, X0, 1.86, 11.6, 0.34,
    [("笔记熟练度，看不见、摸不着 —— 能测吗？", 13, BODY, False, EA, 1.4, 0.3)])
card(s, X0, 2.46, 11.9, 1.40)
txt(s, X0 + 0.34, 2.72, 11.3, 1.0, [
    ("给一个可操作的定义，再编一份可被复核的自评问卷（量表）。", 14, BLACK, False, EA, 1.5, 0.3),
    ("目标不是“发现规律”，而是交付一件能接着用的工具。", 11.5, MUTED, False, EA, 1.5, None)])
txt(s, X0, 4.22, 11.9, 0.6, [
    ("周金华、董燕萍（2019）. 口译笔记熟练度量表的开发. 外语教学与研究 51(6): 925-937.（附 21 题全文）",
     10.5, FAINT, False, EA, 1.5, None)])
txt(s, X0, 5.00, 11.6, 0.32,
    [("构念：两阶段（记录 / 理解）× 四维度（听记协调 · 系统性 · 时效性 · 使用）", 12, MUTED, False, EA, 1.4, None)])
foot(s, 2)

# ============================================================ 03 正例流程
s = S()
tag(s, "正例 · 怎么做", "Procedure")
title(s, "七步，每步留下一件证据", y=1.02)
txt(s, X0, 1.86, 11.4, 0.32, [("右栏 ＝ 你写“研究方法”一节时要写的句子。", 12, MUTED, False, EA, 1.4, None)])
row = [("① 界定构念", "两阶段 × 四维度", "构念定义＋维度表"),
       ("② 生成题库", "74 题，来自日志与访谈", "题项来源说明"),
       ("③ 专家评审", "3 轮 → 28 题（6 点）", "评审人数与依据"),
       ("④ 试测删题", "54 人 → 删 7 题", "删题指标与阈值"),
       ("⑤ 正式施测", "144 人", "总体口径与份数"),
       ("⑥ 信度效度", "α = .901；CVI = 1", "信效度指标表"),
       ("⑦ 三项效标", ".473 · .539 · .556", "效标＋组间比较")]
y = 2.42
for k, a, b in row:
    txt(s, X0, y, 1.62, 0.30, [(k, 12, BLACK, False, EA, 1.3, 0.2)])
    txt(s, 2.42, y, 6.0, 0.30, [(a, 11.5, MUTED, False, EA, 1.3, None)])
    txt(s, 8.62, y, 3.9, 0.30, [(b, 11.5, BODY, False, EA, 1.3, None)])
    line(s, X0, y + 0.30, 11.9, "F0F2F5", 0.6)
    y += 0.52
txt(s, X0, 6.14, 11.6, 0.3,
    [("结果：21 题 · 4 因子解释 60.49% ｜ 初级组与高级组差异显著", 11.5, MUTED, False, EA, 1.4, None)])
foot(s, 3)

# ============================================================ 04 正例三点
s = S()
tag(s, "正例 · 可迁移", "Takeaways")
title(s, "三点值得照搬", y=1.02)
takes = [("构念先于题项", "先有操作定义与对齐表，再写题项。", "Construct first, items second."),
         ("删题是设计，不是失败", "74 → 28 → 21，每轮都留下依据。", "Item analysis, reported."),
         ("结论只走一步", "只说“可测、能区分阶段”，不说“笔记好所以口译好”。", "Claim no more than data.")]
y = 2.44
for i, (a, b, c) in enumerate(takes):
    card(s, X0, y, 11.9, 1.24)
    txt(s, X0 + 0.36, y + 0.24, 11.2, 0.95, [
        ("%d. %s" % (i + 1, a), 15.5, BLACK, True, EA, 1.34, 0.3),
        (b, 11.8, MUTED, False, EA, 1.5, None),
        (c, 10, FAINT, False, LAT, 1.4, None)])
    y += 1.42
txt(s, X0, 6.68, 11.6, 0.26, [("另：21 题随附录公开 —— 工具可复用，研究才能被复核。", 10.5, FAINT, False, EA, 1.4, None)])
foot(s, 4)

# ============================================================ 05 反例概况
s = S()
tag(s, "反例 · 笔译研究", "Negative Case")
title(s, "六道题，去撑一个新流程", y=1.02)
txt(s, X0, 1.86, 11.6, 0.34,
    [("问题：本科翻译专业要不要引入“机器翻译 ＋ 译后编辑（MTPE）”的流程？", 13, BODY, False, EA, 1.4, 0.3)])
card(s, X0, 2.46, 11.9, 1.40)
txt(s, X0 + 0.34, 2.72, 11.3, 1.0, [
    ("工具：6 条李克特陈述（有用性 3 · 易用性 3，6 点）＋ 4 道开放题。", 14, BLACK, False, EA, 1.5, 0.3),
    ("设计：横断面自陈态度；框架借技术接受模型（TAM）；与两门课成绩做相关。", 11.5, MUTED, False, EA, 1.5, None)])
txt(s, X0, 4.22, 11.9, 0.6, [
    ("Yang & Mustafa (2022). On postediting of MT and workflow. HBET, 5793054.（开放获取）",
     10.5, FAINT, False, EA, 1.5, None)])
txt(s, X0, 5.00, 11.6, 0.32,
    [("它也报了数字：N = 127 ｜ α = .714 / .835 ｜ 均值 4.02–4.55（都高于中点）", 12, MUTED, False, EA, 1.4, None)])
foot(s, 5)

# ============================================================ 06 反例：断在哪
s = S()
tag(s, "反例 · 断点", "Where It Breaks")
title(s, "六步都做了，每步只半句", y=1.02)
txt(s, X0, 1.86, 11.6, 0.32, [("左：论文写了的 ｜ 右：缺掉的那一环", 12, MUTED, False, EA, 1.4, None)])
brk = [("① 目的与工具错位", "要论证“可行”，只答了态度", "可行性需评审或试点"),
       ("② 借框架不借结构", "TAM 两构念，6 条自拟题", "结构未经检验"),
       ("③ 抽样口径模糊", "一所高校，便利抽样", "未报发放 / 回收 / 有效"),
       ("④ 施测一次成型", "线上一次填答", "无照应题与偏差控制"),
       ("⑤ 分析层级错位", "单题 × 课程成绩", "把题项当构念用"),
       ("⑥ 写作代替解释", "均值高于中点 → “可行”", "横断面推不出因果")]
y = 2.44
for k, a, b in brk:
    txt(s, X0, y, 2.12, 0.30, [(k, 12, BLACK, False, EA, 1.3, 0.2)])
    txt(s, 3.00, y, 4.32, 0.30, [(a, 11.5, MUTED, False, EA, 1.3, None)])
    txt(s, 7.62, y, 4.9, 0.30, [(b, 11.5, BODY, False, EA, 1.3, None)])
    line(s, X0, y + 0.30, 11.9, "F0F2F5", 0.6)
    y += 0.58
txt(s, X0, 6.06, 11.9, 0.66, [
    ("这不是造假：数据是真的，统计也报了 —— 它只是撑不起被要求撑的那个结论。", 12.5, BLACK, True, EA, 1.5, 0.3)])
foot(s, 6)

# ============================================================ 07 反例三点
s = S()
tag(s, "反例 · 引以为戒", "Lessons")
title(s, "三处最容易踩", y=1.02)
bad = [("题项要来自已验证的量表", "自拟 3 条不等于“感知有用性”。",
        "改：沿用已验证条目并标出处；补 EFA 或报题总相关。"),
       ("别用单题代表构念", "Q6 × 成绩 r = .746 → “能力强”？",
        "改：先合成维度均分并报 α，再做相关；只写“关联”。"),
       ("结论要收缩到样本", "127 人 → “教学可行”？",
        "改：称“该校学生”；报发放 / 回收 / 有效三份数字。")]
y = 2.44
for a, b, c in bad:
    card(s, X0, y, 11.9, 1.24)
    txt(s, X0 + 0.36, y + 0.22, 11.2, 0.95, [
        (a, 15.5, BLACK, True, EA, 1.34, 0.3),
        (b, 11.8, MUTED, False, EA, 1.5, None),
        (c, 11.8, BODY, False, EA, 1.5, None)])
    y += 1.42
txt(s, X0, 6.68, 11.6, 0.26, [("底线：剔卷标准与编码规则要预先写清（此处 18 份被剔除，依据未披露）。",
                               10.5, FAINT, False, EA, 1.4, None)])
foot(s, 7)

# ============================================================ 08 并排对照
s = S()
tag(s, "同一把尺", "Side by Side")
title(s, "两案并排", y=1.02)
txt(s, 6.86, 1.98, 2.4, 0.26, [("口译 · 正例", 11, BLACK, False, EA, 1.3, 0.2)])
txt(s, 9.46, 1.98, 3.06, 0.26, [("笔译 MTPE · 反例", 11, MUTED, False, EA, 1.3, 0.2)])
cmp = [("题项哪来的", "日志＋访谈＋评审 74→28→21", "自拟 6 条，无评审前测"),
       ("测的是什么", "4 维度，EFA 复核吻合", "2 构念各 3 题，未检验"),
       ("谁在填", "144 人，条件写明", "127 人，回收口径未报"),
       ("信度", "α = .901", ".714 / .835"),
       ("效度", "CVI ＋ EFA ＋ 三项效标", "单题 × 课程成绩"),
       ("结论走到哪", "可测、能区分训练阶段", "流程可行、纳入课程")]
y = 2.42
for k, a, b in cmp:
    txt(s, X0, y, 2.1, 0.30, [(k, 11.5, MUTED, False, EA, 1.3, 0.2)])
    txt(s, 6.86, y, 2.42, 0.30, [(a, 11.3, BLACK, False, EA, 1.3, None)])
    txt(s, 9.46, y, 3.06, 0.30 if False else 0.30, [(b, 11.3, MUTED, False, EA, 1.3, None)])
    line(s, X0, y + 0.30, 11.9, "F0F2F5", 0.6)
    y += 0.52
txt(s, X0, 5.66, 5.9, 0.62, [("倒过来读，就是你的自检表", 13, BLACK, True, EA, 1.4, 0.4)])
txt(s, 6.86, 5.72, 5.66, 0.6,
    [("题项哪来的？测的还是它吗？谁填的、留下多少？信度在哪一行？结论多走了几步？",
      11, MUTED, False, EA, 1.5, None)])
foot(s, 8)

# ============================================================ 09 清单
s = S()
tag(s, "交稿前", "Ten Checks")
title(s, "十行自检", y=1.02)
txt(s, X0, 1.86, 11.4, 0.32, [("这十条在正例那篇里都有对应段落 —— 抄它的结构，别抄它的内容。", 12, MUTED, False, EA, 1.4, None)])
ck = ["研究问题：一句可回答的话（描述 / 比较 / 关联）",
      "构念 → 维度 → 题项：一张对齐表",
      "题项来源：沿用注出处｜改编列改动｜自编说过程",
      "总体、抽样框、纳入与排除：各一句话",
      "发放 / 回收 / 有效：三份数字＋回收率",
      "无效判定规则：先定后用",
      "前测与专家评审：留痕",
      "信度 α / ω；多维则 EFA / CFA",
      "指导语：匿名、用途、自愿退出；敏感题后置",
      "结论分寸：横断面只说关联，不说因果"]
y = 2.46
for i, c in enumerate(ck):
    txt(s, X0, y, 0.42, 0.28, [("%02d" % (i + 1), 10, FAINT, False, MONO, 1.3, None)])
    txt(s, 1.14, y, 10.6, 0.28, [(c, 12, BODY, False, EA, 1.3, None)])
    y += 0.42
txt(s, 12.0, 2.46, 0.5, 4.3, [("✓", 12, "D8DCE3", False, EA, 1.3, None)])
foot(s, 9)

# ============================================================ 10 思考题 + 资料
s = S()
tag(s, "讨论 ＋ 资料", "Discussion & Sources")
title(s, "四道思考题", y=1.02)
qs = [("取舍", "30 份问卷的预算，四道检验保哪两道？"),
      ("补一个统计量", "“均值高于中点”＝“总体积极”？你补哪个、为什么？"),
      ("追问样本", "要推广到 MTI，最少补哪三条信息？"),
      ("反身", "对照六病，你最怕踩哪一条？最小改法是什么？")]
y = 2.32
for i, (k, q) in enumerate(qs):
    txt(s, X0, y, 0.9, 0.30, [("Q%d" % (i + 1), 12, NAVY, False, MONO, 1.3, 0.3)])
    txt(s, 1.34, y, 1.5, 0.30, [(k, 12, BLACK, False, EA, 1.3, 0.2)])
    txt(s, 3.02, y, 4.28, 0.30, [(q, 11.5, MUTED, False, EA, 1.34, None)])
    line(s, X0, y + 0.34, 7.3, "F0F2F5", 0.6)
    y += 0.60
txt(s, X0, 4.94, 7.3, 0.5,
    [("讨论 4 分钟：本组每题先自答一句，再点人。", 10.5, FAINT, False, EA, 1.5, None)])
txt(s, 8.34, 1.94, 4.2, 0.3, [("论文与学习资料", 11, BLACK, False, EA, 1.3, 0.3)])
refs = [("正例全文＋量表附录", "bcdlab.gdufs.edu.cn/info/1025/2025.htm"),
        ("反例原文（开放获取）", "onlinelibrary.wiley.com/doi/10.1155/2022/5793054"),
        ("同类笔译态度调查", "doi.org/10.18298/ijlet.3242"),
        ("口译问卷范例", "张威（2013）《中国翻译》(2): 17-25"),
        ("方法读物", "吴明隆《问卷统计分析实务》"),
        ("检索式", "CNKI: 翻译 AND（问卷 OR 量表）AND 信效度")]
y = 2.34
for a, b in refs:
    txt(s, 8.34, y, 4.2, 0.28, [(a, 10.2, BODY, False, EA, 1.32, None)])
    txt(s, 8.34, y + 0.25, 4.2, 0.30, [(b, 8.6, FAINT, False, MONO, 1.25, None)])
    y += 0.66
txt(s, 8.34, y + 0.06, 4.2, 0.4, [("两条链接都可直接点开核对数字；检索式可复制到知网 / Web of Science。", 9.5, FAINT, False, EA, 1.4, None)])
foot(s, 10)

# ---------------------------------------------------------------- 纯白化
prs.save(OUT)
prs2 = Presentation(OUT)
removed = 0
for sl in prs2.slides:
    for sh in list(sl.shapes):
        if sh._element.get("_keep") == "1":
            continue
        try:
            if sh.has_text_frame and sh.text_frame.text.strip():
                continue
        except Exception:
            pass
        sh._element.getparent().remove(sh._element)
        removed += 1
prs2.save(OUT)
tot = 0
for i, sl in enumerate(prs2.slides):
    n = sum(len(sh.text_frame.text.replace("\n", "").replace(" ", "")) for sh in sl.shapes
            if sh.has_text_frame and sh.text_frame.text.strip())
    tot += n
    print("  P%02d 字数 %3d" % (i, n))
print("saved:", OUT, "| slides:", len(list(prs2.slides)), "| decor removed:", removed, "| 全稿字数:", tot)
