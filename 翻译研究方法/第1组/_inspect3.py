# -*- coding: utf-8 -*-
import zipfile, re, sys
from pptx import Presentation

def theme_colors(path):
    z = zipfile.ZipFile(path)
    names = [n for n in z.namelist() if n.endswith('theme1.xml')]
    cols = []
    fonts = set()
    for nm in names:
        t = z.read(nm).decode('utf-8')
        m = re.search(r'<a:clrScheme.*?</a:clrScheme>', t, re.S)
        if m:
            for c in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', m.group(0)):
                cols.append(c)
        for f in re.findall(r'typeface="([^"]+)"', t):
            fonts.add(f)
    return cols, fonts

for path in sys.argv[1:]:
    prs = Presentation(path)
    w = prs.slide_width / 914400
    h = prs.slide_height / 914400
    cols, fonts = theme_colors(path)
    print("=" * 80)
    print("FILE:", path.split('/')[-1])
    print("  size: %.2f x %.2f in | slides: %d" % (w, h, len(prs.slides._sldIdLst)))
    print("  theme accent colors:", cols[:12])
    print("  fonts:", sorted(fonts)[:20])
    try:
        sl = prs.slides[0]
        fcolors = []
        for sh in sl.shapes:
            try:
                c = sh.fill.fore_color
                if c.type is not None:
                    fcolors.append(str(c.rgb))
            except Exception:
                pass
        print("  slide1 fill colors:", fcolors[:20])
    except Exception as e:
        print("  slide1 err", e)
