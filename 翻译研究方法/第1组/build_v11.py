# -*- coding: utf-8 -*-
"""
V11 构建脚本：V9 基础上
- 截图标注精简（generate_underlines.py 已处理）
- 内容补充概念解释（Cronbach's α、自拟题、未报告）
- 新增流程图：正例研究流程（flow）/ 反例证据链断裂（flow_break）
内容源：deck_content_v11.json（30 页）
用法：
    python build_v11.py                # 默认 deck_content_v11.json
    python build_v11.py 润色后.json    # 用润色后的 JSON
"""
import sys, os, json, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

LAT = "Times New Roman"
CN_HEAD = "黑体"
CN_BODY = "宋体"
BLACK = "000000"
GREY = "404040"
BLUE = "5B9BD5"
TEAL = "0E7C86"
TEAL_LT = "E7F1F1"
WARN = "B4451F"
WARN_LT = "F7EBE5"

BASE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "deck_content_v11.json")
OUT = os.path.join(BASE, "问卷调查法正反例_v11.pptx")
ASD = os.path.join(BASE, "assets")
TPL = r"d:\10_Workspace\translation-research-202610\翻译研究方法\template\高级翻译学院专用PPT模板修改版.pptx"

BODY_SIZE = 20
TITLE_SIZE = 23

with open(JSON_PATH, encoding="utf-8") as f:
    data = json.load(f)
slides = data.get("slides", [])

W, H = 13.3333, 7.5
Y_HEAD = 1.25
Y_BODY = 2.05
Y_BOTTOM = 6.35
Y_PAGENO = 6.28

prs = Presentation(TPL)
blank_layout = prs.slides[0].slide_layout
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst):
    rId = sid.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sid)


def set_run(r, text, size, bold=False, cn=CN_BODY, color=BLACK, italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = LAT
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", cn)


def add_tb(slide, x, y, w, h, anchor="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bot": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, text, size, bold=False, cn=CN_BODY, color=BLACK, align="l", first=False,
         lh=1.3, before=0, after=0, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
    p.line_spacing = lh
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    r = p.add_run()
    set_run(r, text, size, bold, cn, color, italic)
    return p


def has_number_prefix(text):
    return re.match(r'^\d+\s*[\.、）)]', text) is not None


def rounded(slide, x, y, w, h, fill, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = 0.12
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def arrow(slide, x, y, w, h, color=TEAL):
    sp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(color)
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def pic_ratio(name):
    path = os.path.join(ASD, name)
    if not os.path.exists(path):
        return 1.5
    im = Image.open(path)
    return im.width / im.height


def contain(w, h, max_w, max_h):
    ratio = w / h
    if ratio > max_w / max_h:
        return max_w, max_w / ratio
    return max_h * ratio, max_h


def add_picture(slide, name, x, y, w, h, caption=None):
    path = os.path.join(ASD, name)
    if not os.path.exists(path):
        return
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - 0.02), Inches(y - 0.02), Inches(w + 0.04), Inches(h + 0.04))
    ln.fill.background()
    ln.line.color.rgb = RGBColor.from_string("BFBFBF")
    ln.line.width = Pt(0.75)
    ln.shadow.inherit = False
    if caption:
        tb, tf = add_tb(slide, x, y + h + 0.05, w, 0.28)
        para(tf, caption, 9.5, cn=CN_BODY, color=GREY, align="c", first=True)


def foot(slide, page, total):
    tb, tf = add_tb(slide, 11.3, Y_PAGENO, 1.8, 0.22)
    para(tf, "%d / %d" % (page, total), 10, cn=CN_BODY, color=GREY, align="r", first=True)


def section_page(s, num, title):
    for bx in (0.0, 9.76):
        blk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(2.5), Inches(3.57), Inches(2.9))
        blk.fill.solid()
        blk.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        blk.line.fill.background()
        blk.shadow.inherit = False
    tb, tf = add_tb(s, 0.9, 2.95, 2.6, 1.6)
    para(tf, num, 96, bold=False, cn=CN_BODY, color="FFFFFF", first=True)
    tb, tf = add_tb(s, 3.9, 3.35, 5.8, 1.4)
    para(tf, title, 28, bold=True, cn=CN_HEAD, color=BLACK, first=True, lh=1.15)


def flow_page(s, title, steps):
    """横向流程：n 个卡片 + 箭头"""
    tb, tf = add_tb(s, 0.7, Y_HEAD, 11.9, 0.7)
    para(tf, title, TITLE_SIZE, bold=True, cn=CN_HEAD, first=True, lh=1.1)
    n = len(steps)
    total_w = 11.9
    gap = 0.32
    card_w = (total_w - gap * (n - 1)) / n
    y = 2.6
    h = 2.4
    for i, step in enumerate(steps):
        x = 0.72 + i * (card_w + gap)
        rounded(s, x, y, card_w, h, TEAL_LT, TEAL, 1.2)
        parts = step.split("\n")
        tb, tf = add_tb(s, x + 0.1, y + 0.5, card_w - 0.2, h - 0.9, anchor="mid")
        para(tf, parts[0], 15, bold=True, cn=CN_HEAD, color=BLACK, align="c", first=True, lh=1.1)
        if len(parts) > 1:
            para(tf, parts[1], 11, cn=CN_BODY, color=GREY, align="c", lh=1.15)
        if i < n - 1:
            arrow(s, x + card_w + 0.03, y + h / 2 - 0.13, gap - 0.06, 0.26)
    foot(s, 0, 0)


def flow_break_page(s, title, steps):
    """反例证据链断裂：环节卡片 + 下方红色问题标注"""
    tb, tf = add_tb(s, 0.7, Y_HEAD, 11.9, 0.7)
    para(tf, title, TITLE_SIZE, bold=True, cn=CN_HEAD, first=True, lh=1.1)
    n = len(steps)
    total_w = 11.9
    gap = 0.3
    card_w = (total_w - gap * (n - 1)) / n
    y = 2.5
    h = 1.4
    for i, st in enumerate(steps):
        x = 0.72 + i * (card_w + gap)
        # 环节卡片
        rounded(s, x, y, card_w, h, "FFFFFF", "9AA0A6", 1.2)
        tb, tf = add_tb(s, x + 0.08, y, card_w - 0.16, h, anchor="mid")
        para(tf, st["name"], 15, bold=True, cn=CN_HEAD, color=BLACK, align="c", first=True, lh=1.1)
        # 红色问题标注
        rounded(s, x, y + h + 0.35, card_w, 0.85, WARN_LT, WARN, 1.0)
        tb, tf = add_tb(s, x + 0.08, y + h + 0.35, card_w - 0.16, 0.85, anchor="mid")
        para(tf, "✕ " + st["issue"], 12.5, bold=True, cn=CN_BODY, color=WARN, align="c", first=True, lh=1.15)
        if i < n - 1:
            arrow(s, x + card_w + 0.02, y + h / 2 - 0.13, gap - 0.04, 0.26, "9AA0A6")
    foot(s, 0, 0)


total = len(slides)

for sd in slides:
    typ = sd.get("type", "slide")
    page = sd.get("page", 0)
    s = prs.slides.add_slide(blank_layout)

    if typ == "cover":
        tb, tf = add_tb(s, 1.0, 2.5, 11.3, 1.0)
        para(tf, sd.get("title", ""), 36, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.7, 11.3, 0.5)
            para(tf, sd.get("subtitle", ""), 20, cn=CN_BODY, align="c", first=True)
        if sd.get("subtitle_en"):
            tb, tf = add_tb(s, 1.0, 4.3, 11.3, 0.4)
            para(tf, sd.get("subtitle_en", ""), 15, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 5.3, 11.3, 1.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 14, cn=CN_BODY, align="c", first=(i == 0), lh=1.4)

    elif typ == "toc":
        tb, tf = add_tb(s, 0.7, Y_HEAD, 6.0, 0.6)
        para(tf, sd.get("title", "目录"), 26, bold=True, cn=CN_HEAD, first=True)
        tb, tf = add_tb(s, 0.7, Y_BODY + 0.1, 11.9, 4.2)
        for i, it in enumerate(sd.get("items", [])):
            para(tf, it, 18, cn=CN_BODY, first=(i == 0), lh=1.3, after=12)
        foot(s, page, total)

    elif typ == "section":
        section_page(s, sd.get("number", ""), sd.get("title", ""))
        foot(s, page, total)

    elif typ == "flow":
        flow_page(s, sd.get("title", ""), sd.get("steps", []))
        foot(s, page, total)

    elif typ == "flow_break":
        flow_break_page(s, sd.get("title", ""), sd.get("steps", []))
        foot(s, page, total)

    elif typ == "slide":
        tb, tf = add_tb(s, 0.7, Y_HEAD, 11.9, 0.7)
        para(tf, sd.get("title", ""), TITLE_SIZE, bold=True, cn=CN_HEAD, first=True, lh=1.1)
        img = sd.get("image")
        img2 = sd.get("image2")

        if img2:
            max_w, max_h = 4.6, 2.05
            p1 = os.path.join(ASD, img)
            p2 = os.path.join(ASD, img2)
            sz1 = Image.open(p1).size if os.path.exists(p1) else (1, 1)
            sz2 = Image.open(p2).size if os.path.exists(p2) else (1, 1)
            w1, h1 = contain(sz1[0], sz1[1], max_w, max_h)
            w2, h2 = contain(sz2[0], sz2[1], max_w, max_h)
            x1 = 12.85 - w1
            x2 = 12.85 - w2
            y1 = 1.9
            y2 = y1 + h1 + 0.35
            lw = min(x1, x2) - 0.9
            bottom = Y_BOTTOM
            add_picture(s, img, x1, y1, w1, h1, sd.get("caption"))
            add_picture(s, img2, x2, y2, w2, h2, sd.get("caption2"))
        elif img:
            ratio = pic_ratio(img)
            if ratio > 2.8:
                h = 1.9
                w = min(1.9 * ratio, 9.5)
                h = w / ratio
                x = (W - w) / 2
                y = 4.3
                lw = 11.9
                bottom = 4.2
            else:
                p = os.path.join(ASD, img)
                sz = Image.open(p).size if os.path.exists(p) else (1, 1)
                w, h = contain(sz[0], sz[1], 4.6, 4.0)
                x = 12.85 - w
                y = 1.95
                lw = x - 0.9
                bottom = Y_BOTTOM
            add_picture(s, img, x, y, w, h, sd.get("caption"))
        else:
            lw = 11.9
            bottom = Y_BOTTOM

        tb, tf = add_tb(s, 0.7, Y_BODY, lw, bottom - Y_BODY)
        for i, b in enumerate(sd.get("bullets", [])):
            txt = b if has_number_prefix(b) else "· " + b
            para(tf, txt, BODY_SIZE, cn=CN_BODY, first=(i == 0), lh=1.3, after=5)
        foot(s, page, total)

    elif typ == "end":
        tb, tf = add_tb(s, 1.0, 2.4, 11.3, 1.0)
        para(tf, sd.get("title", "谢谢观看"), 40, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.5, 11.3, 0.4)
            para(tf, sd.get("subtitle", ""), 18, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 4.1, 11.3, 2.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 12, cn=CN_BODY, align="c", first=(i == 0), lh=1.5)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(slides), "| 正文字号:", BODY_SIZE, "| 内容源:", os.path.basename(JSON_PATH))
