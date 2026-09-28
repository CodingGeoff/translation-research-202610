# -*- coding: utf-8 -*-
"""
V8 构建脚本：V7（模板 + logo + 格言 + 蓝色章节页）基础上
- 正文放大到 20 号字（原先 14.5）
- 论文截图已由 generate_highlights.py 换成「精确定位 + 高亮 + 裁剪」版本
用法：
    python build_v8.py                # 默认 deck_content_v5.json
    python build_v8.py 润色后.json    # 用润色后的 JSON
"""
import sys, os, json, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

LAT = "Times New Roman"
CN_HEAD = "黑体"
CN_BODY = "宋体"
BLACK = "000000"
GREY = "404040"
BLUE = "5B9BD5"

BASE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "deck_content_v5.json")
OUT = os.path.join(BASE, "问卷调查法正反例_v8.pptx")
ASD = os.path.join(BASE, "assets")
TPL = r"d:\10_Workspace\translation-research-202610\翻译研究方法\template\高级翻译学院专用PPT模板修改版.pptx"

BODY_SIZE = 20          # 正文字号（V8 放大）
TITLE_SIZE = 23         # 标题字号

with open(JSON_PATH, encoding="utf-8") as f:
    data = json.load(f)
slides = data.get("slides", [])

W, H = 13.3333, 7.5
Y_HEAD = 1.25
Y_BODY = 2.05
Y_BOTTOM = 6.35
Y_PAGENO = 6.28
IMG_H = 3.9           # 右侧图片高度
IMG_RIGHT = 12.85     # 图片右边缘

prs = Presentation(TPL)
blank_layout = prs.slides[0].slide_layout
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst):
    rId = sid.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sid)


def set_run(r, text, size, bold=False, cn=CN_BODY, color=BLACK, italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = LAT
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", cn)


def add_tb(slide, x, y, w, h, anchor="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bot": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, text, size, bold=False, cn=CN_BODY, color=BLACK, align="l", first=False,
         lh=1.3, before=0, after=0, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
    p.line_spacing = lh
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    r = p.add_run()
    set_run(r, text, size, bold, cn, color, italic)
    return p


def has_number_prefix(text):
    return re.match(r'^\d+\s*[\.、）)]', text) is not None


def img_layout(name):
    """根据截图宽高比决定布局：横版放底部，方形/竖版放右侧。
    返回 (x, y, w, h, narrow) —— narrow=True 表示文字用窄栏。"""
    path = os.path.join(ASD, name)
    if not os.path.exists(path):
        return 9.0, 2.0, 4.0, 3.0, True
    im = Image.open(path)
    ratio = im.width / im.height
    if ratio > 2.8:
        # 横版：底部居中，高 1.9（宽度按比例，不超过 9.5）
        h = 1.9
        w = 1.9 * ratio
        if w > 9.5:
            w = 9.5
            h = 9.5 / ratio
        return (W - w) / 2, 4.3, w, h, False
    # 方形/竖版：右侧 contain（限宽 4.6、限高 4.0）
    max_w, max_h = 4.6, 4.0
    if ratio > max_w / max_h:
        w = max_w
        h = max_w / ratio
    else:
        h = max_h
        w = max_h * ratio
    return 12.85 - w, 1.95, w, h, True


def add_picture(slide, name, x, y, w, h, caption=None):
    path = os.path.join(ASD, name)
    if not os.path.exists(path):
        return
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - 0.02), Inches(y - 0.02), Inches(w + 0.04), Inches(h + 0.04))
    ln.fill.background()
    ln.line.color.rgb = RGBColor.from_string("BFBFBF")
    ln.line.width = Pt(0.75)
    ln.shadow.inherit = False
    if caption:
        tb, tf = add_tb(slide, x, y + h + 0.06, w, 0.3)
        para(tf, caption, 10, cn=CN_BODY, color=GREY, align="c", first=True)


def foot(slide, page, total):
    tb, tf = add_tb(slide, 11.3, Y_PAGENO, 1.8, 0.22)
    para(tf, "%d / %d" % (page, total), 10, cn=CN_BODY, color=GREY, align="r", first=True)


def section_page(s, num, title):
    for bx in (0.0, 9.76):
        blk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(2.5), Inches(3.57), Inches(2.9))
        blk.fill.solid()
        blk.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        blk.line.fill.background()
        blk.shadow.inherit = False
    tb, tf = add_tb(s, 0.9, 2.95, 2.6, 1.6)
    para(tf, num, 96, bold=False, cn=CN_BODY, color="FFFFFF", first=True)
    tb, tf = add_tb(s, 3.9, 3.35, 5.8, 1.4)
    para(tf, title, 28, bold=True, cn=CN_HEAD, color=BLACK, first=True, lh=1.15)


total = len(slides)

for sd in slides:
    typ = sd.get("type", "slide")
    page = sd.get("page", 0)
    s = prs.slides.add_slide(blank_layout)

    if typ == "cover":
        tb, tf = add_tb(s, 1.0, 2.5, 11.3, 1.0)
        para(tf, sd.get("title", ""), 36, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.7, 11.3, 0.5)
            para(tf, sd.get("subtitle", ""), 20, cn=CN_BODY, align="c", first=True)
        if sd.get("subtitle_en"):
            tb, tf = add_tb(s, 1.0, 4.3, 11.3, 0.4)
            para(tf, sd.get("subtitle_en", ""), 15, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 5.3, 11.3, 1.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 14, cn=CN_BODY, align="c", first=(i == 0), lh=1.4)

    elif typ == "toc":
        tb, tf = add_tb(s, 0.7, Y_HEAD, 6.0, 0.6)
        para(tf, sd.get("title", "目录"), 26, bold=True, cn=CN_HEAD, first=True)
        tb, tf = add_tb(s, 0.7, Y_BODY + 0.1, 11.9, 4.2)
        for i, it in enumerate(sd.get("items", [])):
            para(tf, it, 18, cn=CN_BODY, first=(i == 0), lh=1.3, after=12)
        foot(s, page, total)

    elif typ == "section":
        section_page(s, sd.get("number", ""), sd.get("title", ""))
        foot(s, page, total)

    elif typ == "slide":
        tb, tf = add_tb(s, 0.7, Y_HEAD, 11.9, 0.7)
        para(tf, sd.get("title", ""), TITLE_SIZE, bold=True, cn=CN_HEAD, first=True, lh=1.1)
        img = sd.get("image")
        has_img = bool(img)
        if has_img:
            ix, iy, iw, ih, narrow = img_layout(img)
            lw = (ix - 0.9) if narrow else 11.9
        else:
            lw = 11.9
            ix = iy = iw = ih = narrow = None
        # 底部横版图时，文字区要避开图片（文字止于图片上方）
        bottom = (4.2 if (has_img and not narrow) else Y_BOTTOM)
        tb, tf = add_tb(s, 0.7, Y_BODY, lw, bottom - Y_BODY)
        for i, b in enumerate(sd.get("bullets", [])):
            txt = b if has_number_prefix(b) else "· " + b
            para(tf, txt, BODY_SIZE, cn=CN_BODY, first=(i == 0), lh=1.3, after=5)
        if has_img:
            add_picture(s, img, ix, iy, iw, ih, sd.get("caption"))
        foot(s, page, total)

    elif typ == "end":
        tb, tf = add_tb(s, 1.0, 2.4, 11.3, 1.0)
        para(tf, sd.get("title", "谢谢观看"), 40, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.5, 11.3, 0.4)
            para(tf, sd.get("subtitle", ""), 18, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 4.1, 11.3, 2.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 12, cn=CN_BODY, align="c", first=(i == 0), lh=1.5)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(slides), "| 正文字号:", BODY_SIZE, "| 内容源:", os.path.basename(JSON_PATH))
