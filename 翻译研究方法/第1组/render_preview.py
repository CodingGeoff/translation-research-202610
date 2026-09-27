# -*- coding: utf-8 -*-
"""离线渲染预览 + 溢出检查（用 PIL 近似重绘 PPTX，用于人工核对版面）"""
import os, sys, math
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

EMU_IN = 914400.0
SCALE = 150  # px per inch
F_REG = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
F_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
IDX = 2
_fc = {}


def font(px, bold):
    k = (px, bold)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(F_BOLD if bold else F_REG, max(6, int(round(px))), index=IDX)
    return _fc[k]


def emu2px(v):
    return float(v) / EMU_IN * SCALE


def norm(v):
    if v is None: return None
    t = str(v)
    if len(t)==6 and all(c in "0123456789abcdefABCDEF" for c in t):
        return "#"+t
    return t if t.startswith("#") else "#FFFFFF"


def solid_hex(sp):
    try:
        if sp.fill.type is not None and str(sp.fill.type) == "MSO_FILL_TYPE.SOLID" or sp.fill.type == 1:
            c = sp.fill.fore_color
            if c and c.type is not None and str(c.type).startswith("MSO_THEME"):
                return "#FFFFFF"
            return norm(str(c.rgb))
    except Exception:
        pass
    return None


def lines_of(shape, w_px):
    out = []
    try:
        tf = shape.text_frame
    except Exception:
        return out
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        r0 = p.runs[0]
        sz = (r0.font.size.pt if r0.font.size else 18)
        bold = bool(r0.font.bold)
        col = "#111A2C"
        try:
            if r0.font.color and r0.font.color.type is not None:
                col = norm(str(r0.font.color.rgb))
        except Exception:
            pass
        lh = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        px = sz / 72.0 * SCALE
        out.append((txt, px, bold, col, lh, sz))
    return out


def wrap_pil(txt, px, bold, maxw):
    f = font(px, bold)
    out, cur = [], ""
    for para in txt.split("\n"):
        if para == "":
            out.append("")
            continue
        for ch in para:
            if f.getlength(cur + ch) <= maxw or not cur:
                cur += ch
            else:
                out.append(cur)
                cur = ch
    if cur:
        out.append(cur)
    return out


def rot_size(w, h, deg):
    a = math.radians(deg)
    nw = abs(w * math.cos(a)) + abs(h * math.sin(a))
    nh = abs(w * math.sin(a)) + abs(h * math.cos(a))
    return nw, nh


def render(pptx_path, outdir="preview", max_slides=None):
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(pptx_path)
    W = emu2px(prs.slide_width)
    H = emu2px(prs.slide_height)
    warnings = []
    imgs = []
    for idx, slide in enumerate(prs.slides, start=1):
        if max_slides and idx > max_slides:
            break
        img = Image.new("RGB", (int(W), int(H)), "white")
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            try:
                x, y = emu2px(sh.left), emu2px(sh.top)
                w, h = emu2px(sh.width), emu2px(sh.height)
            except Exception:
                continue
            name = sh.shape_type
            rot = getattr(sh, "rotation", 0) or 0
            fill = solid_hex(sh) if hasattr(sh, "fill") else None
            stroke = None
            try:
                if sh.line.color and sh.line.color.type is not None:
                    stroke = norm(str(sh.line.color.rgb))
            except Exception:
                pass
            auto_shape = ""
            try:
                auto_shape = str(sh.auto_shape_type)
            except Exception:
                pass
            if (x < -8 or y < -8 or x + w > W + 10 or y + h > H + 10) and (w < W - 4 or h < H - 4):
                    warnings.append("S%02d 形状超出画布: %s (%.2f,%.2f,%.2f,%.2f)" %
                                    (idx, auto_shape or name, x / SCALE, y / SCALE, w / SCALE, h / SCALE))
            if fill is not None or stroke is not None:
                if "ROUNDED" in auto_shape:
                    r = min(w, h) * 0.18
                    d.rounded_rectangle([x, y, x + w, y + h], radius=r, fill=fill, outline=stroke, width=2)
                elif "OVAL" in auto_shape:
                    d.ellipse([x, y, x + w, y + h], fill=fill, outline=stroke, width=2)
                elif "PENTAGON" in auto_shape:
                    pts = [(x, y), (x + w * .84, y), (x + w, y + h / 2), (x + w * .84, y + h), (x, y + h)]
                    d.polygon(pts, fill=fill, outline=stroke)
                elif "TRIANGLE" in auto_shape:
                    nw, nh = rot_size(w, h, rot)
                    cx, cy = x + w / 2, y + h / 2
                    pts = [(cx - w / 2, cy + h / 2), (cx + w / 2, cy + h / 2), (cx, cy - h / 2)]
                    d.polygon(pts, fill=fill)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=fill, outline=stroke, width=2)
            # text
            paras = lines_of(sh, w)
            if not paras:
                continue
            pad_l = pad_t = 0.0
            try:
                pad_l = emu2px(sh.text_frame.margin_left or 0)
                pad_t = emu2px(sh.text_frame.margin_top or 0)
            except Exception:
                pass
            tw_ = max(10, w - pad_l * 2)
            yy = y + pad_t
            total_h = 0
            for txt, px, bold, col, lh, ptsz in paras:
                ls = wrap_pil(txt, px, bold, tw_)
                total_h += len(ls) * px * lh
            anchor = "t"
            try:
                va = str(sh.text_frame.vertical_anchor)
                if va and "MIDDLE" in va:
                    anchor = "m"
                elif va and "BOTTOM" in va:
                    anchor = "b"
            except Exception:
                pass
            if anchor == "m" and total_h < h:
                yy = y + (h - total_h) / 2
            elif anchor == "b":
                yy = y + max(0, h - total_h) - pad_t
            for txt, px, bold, col, lh, ptsz in paras:
                ls = wrap_pil(txt, px, bold, tw_)
                for ln in ls:
                    if yy + px * lh > H + 1:
                        warnings.append("S%02d 文本超出页底: “%s…”" % (idx, ln[:16]))
                        break
                    d.text((x + pad_l, yy), ln, font=font(px, bold), fill=col)
                    yy += px * lh
            need_h = total_h
            box_h = h
            if need_h > box_h + 6 and box_h > 8:
                warnings.append("S%02d 文本可能溢出文本框(需 %.2fin / 框 %.2fin)：“%s”" %
                                (idx, need_h / SCALE, box_h / SCALE, paras[0][0][:18]))
        p = os.path.join(outdir, "s%02d.png" % idx)
        img.save(p)
        imgs.append(p)
    return imgs, warnings


def contact(imgs, out, per=9, cols=3, cellw=760):
    n = len(imgs)
    rows = -(-n // per)
    ims = [Image.open(p) for p in imgs]
    if not ims:
        return
    ratio = ims[0].height / ims[0].width
    cw = cellw
    ch = int(cw * ratio)
    pad = 14
    sheet_h = rows * ((per // cols) * (ch + pad) + pad) + pad
    sheets = 0
    for r in range(rows):
        block = imgs[r * per:(r + 1) * per]
        canvas = Image.new("RGB", (cols * cw + (cols + 1) * pad,
                                   -(-len(block) // cols) * (ch + pad) + pad * 2), "#E8E6E0")
        cd = ImageDraw.Draw(canvas)
        for i, p in enumerate(block):
            cx = pad + (i % cols) * (cw + pad)
            cy = pad + (i // cols) * (ch + pad)
            im = Image.open(p).resize((cw, ch))
            canvas.paste(im, (cx, cy))
            cd.rectangle([cx, cy, cx + cw, cy + ch], outline="#B9B4A8", width=1)
            cd.text((cx + 6, cy + 4), os.path.basename(p), font=font(15, True), fill="#5A564C")
        o = out.replace(".png", "_%d.png" % (sheets + 1))
        canvas.save(o)
        sheets += 1
    return sheets


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "第1组_问卷调查法与访谈法_v1.pptx"
    imgs, warn = render(src)
    print("slides rendered:", len(imgs))
    if warn:
        print("--- warnings (%d) ---" % len(warn))
        for w in warn:
            print(w)
    else:
        print("no layout warnings")
    open("check_report.txt", "w").write("\n".join(warn) if warn else "no layout warnings")
    contact(imgs, "preview/contact.png")
    print("contact sheets written")
