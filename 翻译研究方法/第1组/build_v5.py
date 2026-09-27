# -*- coding: utf-8 -*-
"""
V5 构建脚本：读 deck_content_v5.json 生成 PPT
- 风格：模仿学长学姐答辩 PPT —— 白底黑字、Times New Roman + 黑体/宋体
- 内容源与生成分离：改 JSON 文字即可重新生成（可交给别的 AI 润色 JSON）
用法：
    python build_v5.py                # 用默认 deck_content_v5.json 生成
    python build_v5.py 润色后.json    # 用润色后的 JSON 生成
"""
import sys, os, json, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

# ---- 字体与颜色 ----
LAT = "Times New Roman"      # 西文/数字
CN_HEAD = "黑体"              # 标题中文
CN_BODY = "宋体"              # 正文中文
BLACK = "000000"
GREY = "404040"               # 次要文字仍用深灰，接近黑

BASE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "deck_content_v5.json")
OUT = os.path.join(BASE, "问卷调查法正反例_v5.pptx")
ASD = os.path.join(BASE, "assets")

with open(JSON_PATH, encoding="utf-8") as f:
    data = json.load(f)

meta = data.get("meta", {})
slides = data.get("slides", [])

W, H = 13.3333, 7.5
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def set_run(r, text, size, bold=False, cn=CN_BODY, color=BLACK, italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = LAT
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", cn)


def add_tb(slide, x, y, w, h, anchor="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bot": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, text, size, bold=False, cn=CN_BODY, color=BLACK, align="l", first=False,
         lh=1.3, before=0, after=0, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
    p.line_spacing = lh
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    r = p.add_run()
    set_run(r, text, size, bold, cn, color, italic)
    return p


def has_number_prefix(text):
    return re.match(r'^\d+\s*[\.、）)]', text) is not None


def add_picture(slide, name, x, y, h, caption=None):
    path = os.path.join(ASD, name)
    if not os.path.exists(path):
        return
    im = Image.open(path)
    w = h * im.width / im.height
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    # 细边框
    from pptx.enum.shapes import MSO_SHAPE
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - 0.02), Inches(y - 0.02), Inches(w + 0.04), Inches(h + 0.04))
    ln.fill.background()
    ln.line.color.rgb = RGBColor.from_string("BFBFBF")
    ln.line.width = Pt(0.75)
    ln.shadow.inherit = False
    if caption:
        tb, tf = add_tb(slide, x, y + h + 0.06, w, 0.3)
        para(tf, caption, 10, cn=CN_BODY, color=GREY, align="c", first=True)


def foot(slide, page, total):
    tb, tf = add_tb(slide, 11.5, 7.1, 1.6, 0.3)
    para(tf, "%d / %d" % (page, total), 10, cn=CN_BODY, color=GREY, align="r", first=True)


total = len(slides)

for sd in slides:
    typ = sd.get("type", "slide")
    page = sd.get("page", 0)
    s = prs.slides.add_slide(blank)

    if typ == "cover":
        tb, tf = add_tb(s, 1.0, 2.3, 11.3, 1.0)
        para(tf, sd.get("title", ""), 36, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.5, 11.3, 0.5)
            para(tf, sd.get("subtitle", ""), 20, cn=CN_BODY, align="c", first=True)
        if sd.get("subtitle_en"):
            tb, tf = add_tb(s, 1.0, 4.1, 11.3, 0.4)
            para(tf, sd.get("subtitle_en", ""), 15, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 5.6, 11.3, 1.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 14, cn=CN_BODY, align="c", first=(i == 0), lh=1.4)

    elif typ == "toc":
        tb, tf = add_tb(s, 0.7, 0.7, 6.0, 0.7)
        para(tf, sd.get("title", "目录"), 28, bold=True, cn=CN_HEAD, first=True)
        tb, tf = add_tb(s, 0.7, 2.0, 11.9, 4.6)
        for i, it in enumerate(sd.get("items", [])):
            para(tf, it, 18, cn=CN_BODY, first=(i == 0), lh=1.3, after=14)
        foot(s, page, total)

    elif typ == "section":
        num = sd.get("number", "")
        tb, tf = add_tb(s, 0.9, 2.2, 3.0, 1.6)
        para(tf, num, 60, bold=True, cn=CN_BODY, color=GREY, first=True)
        tb, tf = add_tb(s, 4.0, 2.5, 8.5, 1.2)
        para(tf, sd.get("title", ""), 32, bold=True, cn=CN_HEAD, first=True, lh=1.15)
        foot(s, page, total)

    elif typ == "slide":
        tb, tf = add_tb(s, 0.7, 0.55, 11.9, 0.7)
        para(tf, sd.get("title", ""), 24, bold=True, cn=CN_HEAD, first=True, lh=1.1)
        has_img = bool(sd.get("image"))
        lx, lw = 0.7, (6.9 if has_img else 11.9)
        tb, tf = add_tb(s, lx, 1.6, lw, 5.2)
        for i, b in enumerate(sd.get("bullets", [])):
            txt = b if has_number_prefix(b) else "· " + b
            para(tf, txt, 15, cn=CN_BODY, first=(i == 0), lh=1.32, after=10)
        if has_img:
            add_picture(s, sd.get("image"), 8.0, 1.75, 4.3, sd.get("caption"))
        foot(s, page, total)

    elif typ == "end":
        tb, tf = add_tb(s, 1.0, 2.2, 11.3, 1.0)
        para(tf, sd.get("title", "谢谢观看"), 40, bold=True, cn=CN_HEAD, align="c", first=True, lh=1.1)
        if sd.get("subtitle"):
            tb, tf = add_tb(s, 1.0, 3.4, 11.3, 0.4)
            para(tf, sd.get("subtitle", ""), 18, cn=CN_BODY, color=GREY, align="c", first=True)
        tb, tf = add_tb(s, 1.0, 4.2, 11.3, 2.0)
        for i, ln in enumerate(sd.get("lines", [])):
            para(tf, ln, 12.5, cn=CN_BODY, align="c", first=(i == 0), lh=1.5)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(slides), "| 内容源:", os.path.basename(JSON_PATH))
