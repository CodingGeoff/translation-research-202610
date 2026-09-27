# -*- coding: utf-8 -*-
"""
第1组《翻译研究方法》课堂展示 PPT 生成器 —— 版式引擎
问卷调查法 (Questionnaire) + 访谈法 (Interview)
"""
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ----------------------------------------------------------------- 基础常量
W, H = 13.3333, 7.5
ML = 0.62
CR = W - ML
CT, CB = 1.70, 6.86

EA = "Microsoft YaHei"
LAT = "Arial"
MONO = "Consolas"

INK = "111A2C"
INK_SOFT = "24304a"
PAPER = "FCFBF8"
CARD = "FFFFFF"
GREY_T = "6C7480"
LINE = "D9D5CC"
GOOD = "1E7A5C"
WARN = "B4451F"
TEAL = "2B7C85"
C1, C2, C3, C4 = "D2603A", "C8963E", "2B7C85", "6B63A6"

F_REG = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
F_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
IDX = 2  # SC

_cache = {}


def _pil(size, bold=False):
    key = (size, bold)
    if key not in _cache:
        path = F_BOLD if bold else F_REG
        _cache[key] = ImageFont.truetype(path, int(round(size)), index=IDX)
    return _cache[key]


def tw(txt, size, bold=False):
    """文本宽度（英寸），按渲染字号 size(pt) 计算"""
    if not txt:
        return 0.0
    f = _pil(size * 96 / 72.0, bold)
    return f.getlength(txt) / 96.0


def wrap(txt, width, size, bold=False):
    """按可用宽度折行，返回行列表（支持显式换行 \n）"""
    out = []
    for para in str(txt).split("\n"):
        if not para:
            out.append("")
            continue
        cur = ""
        for ch in para:
            trial = cur + ch
            if tw(trial, size, bold) <= width or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = ch
        out.append(cur)
    return out


def fit_wrap(txt, box_w, box_h, size, lh=1.34, bold=False, min_size=9.5):
    """自动缩字号直到文字装进 box_h；返回 (lines, size, used_h)"""
    s = size
    while True:
        lines = wrap(txt, box_w, s, bold)
        used = len(lines) * s * lh / 72.0
        if used <= box_h or s <= min_size:
            return lines, s, used
        s = round(s - 0.35, 2)


def hx(h):
    return RGBColor.from_string(h)


def tint(hexcol, k=0.10, base=PAPER):
    """淡色底：k 为原色占比（越小越浅）"""
    a = [int(hexcol[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(base[i:i + 2], 16) for i in (0, 2, 4)]
    c = [round(b[i] + (a[i] - b[i]) * k) for i in range(3)]
    return "%02X%02X%02X" % tuple(c)


# ----------------------------------------------------------------- 底层绘制
def set_run(r, txt, size=12.5, color=INK, bold=False, font=EA, italic=False, spc=None):
    r.text = txt
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = hx(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", EA)
    if spc is not None:
        rPr.set("spc", str(int(spc * 100)))


def add_tb(slide, x, y, w, h, anchor="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE,
                          "bot": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, txt, size=12.5, color=INK, bold=False, lh=1.32, before=0, after=0,
         align="l", font=EA, first=False, spc=None, indent=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
                   "j": PP_ALIGN.JUSTIFY}[align]
    p.line_spacing = lh
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if indent:
        p.paragraph_properties = getattr(p, "paragraph_properties", None)
    r = p.add_run()
    set_run(r, txt, size, color, bold, font, spc=spc)
    return p


def rect(slide, x, y, w, h, fill=None, ln=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
         adj=None, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = hx(fill)
    if ln is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = hx(ln)
        s.line.width = Pt(lw)
    if adj is not None:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return s


def card(slide, x, y, w, h, fill=CARD, ln=LINE, lw=0.75, top_bar=None, bar_h=0.055,
          radius=None):
    if radius is not None:
        s = rect(slide, x, y, w, h, fill, ln, lw, MSO_SHAPE.ROUNDED_RECTANGLE, adj=radius)
    else:
        s = rect(slide, x, y, w, h, fill, ln, lw)
    if top_bar:
        rect(slide, x, y, w, bar_h, top_bar)
    return s


# ----------------------------------------------------------------- 版式
def bg(slide, color=PAPER):
    rect(slide, -0.06, -0.06, W + 0.12, H + 0.12, color)


def footer(slide, no, total, accent=INK_SOFT, light=False):
    col = "C9CFDA" if light else GREY_T
    tb, tf = add_tb(slide, ML, 6.90, 8.4, 0.26)
    para(tf, "《翻译研究方法》2026–2027 学年第 1 学期 · 第 1 组 · 10.12 课堂展示",
         size=9.5, color=col, first=True)
    tb, tf = add_tb(slide, 9.30, 7.14, 3.41, 0.24)
    p = para(tf, "", size=9.5, color=col, align="r", first=True)
    r = p.add_run()
    set_run(r, "第 %02d / %d 页" % (no, total), 9.5, col, True, MONO)


def segbar(slide, seg_idx, accent):
    """底部四段进度条：标明当前发言人"""
    labels = ["于萍 · 方法 A 概述", "陈冠臻 · 方法 A 案例", "熊芮 · 方法 B 概述", "方燕 · 方法 B 案例"]
    cols = [C1, C1, C3, C3]
    x, w, y, h = 9.30, 3.41, 6.72, 0.15
    seg_w = w / 4.0
    for i in range(4):
        col = accent if i == seg_idx else "E4E1D9"
        rect(slide, x + i * seg_w + (0.0 if i else 0), y, seg_w - 0.035, h, col)
    nm, topic = labels[seg_idx].split(" · ")
    tb, tf = add_tb(slide, x, y + 0.19, w, 0.21)
    para(tf, "主讲 %s · %s" % (nm, topic), size=8.6, color=accent, bold=True,
         align="r", first=True)


def head(slide, eyebrow, title, sub=None, accent=C1, seg=None, no=0, total=35):
    bg(slide)
    rect(slide, ML, 0.55, 0.105, 0.105, accent)
    tb, tf = add_tb(slide, ML + 0.24, 0.50, 9.6, 0.26)
    para(tf, eyebrow, size=10, color=GREY_T, bold=True, first=True, spc=0.9)
    tb, tf = add_tb(slide, ML, 0.82, 11.4, 0.55)
    para(tf, title, size=24, color=INK, bold=True, first=True, lh=1.05)
    if sub:
        tb, tf = add_tb(slide, ML, 1.36, 12.09, 0.26)
        para(tf, sub, size=11, color=GREY_T, first=True, font=MONO)
    rect(slide, ML, 1.63, 12.09, 0.012, LINE)
    if seg is not None:
        segbar(slide, seg, accent)
    footer(slide, no, total)


def dark_head(slide, kicker, title, en=None, accent=C1):
    bg(slide, INK)
    rect(slide, 0, 0, W, 0.055, accent)
    tb, tf = add_tb(slide, 1.05, 1.55, 10, 0.3)
    para(tf, kicker, size=11, color=accent, bold=True, first=True, spc=1.2)
    tb, tf = add_tb(slide, 1.05, 2.00, 11.0, 1.4)
    para(tf, title, size=40, color="F5F3EE", bold=True, first=True, lh=1.12)
    if en:
        tb, tf = add_tb(slide, 1.05, 3.62, 11.0, 0.4)
        para(tf, en, size=13.5, color="9AA4B8", first=True, font=MONO)


# ----------------------------------------------------------------- 内容组件
def card_grid(slide, items, bounds, cols=2, num_bg=None, num_fg="FFFFFF",
              tsize=13.5, bsize=12.0, accent=None, top_bar=True):
    x, y, w, h = bounds
    rows = -(-len(items) // cols)
    gapx, gapy = 0.20, 0.16
    cw = (w - gapx * (cols - 1)) / cols
    ch = (h - gapy * (rows - 1)) / rows
    for i, it in enumerate(items):
        r_, c_ = divmod(i, cols)
        cx, cy = x + c_ * (cw + gapx), y + r_ * (ch + gapy)
        acc = (accent[i % len(accent)] if isinstance(accent, list)
               else (accent or num_bg or INK_SOFT))
        card(slide, cx, cy, cw, ch, CARD, LINE, 0.75, acc if top_bar else None)
        tx = cx + 0.20
        ty = cy + (0.20 if top_bar else 0.16)
        tw_ = cw - 0.40
        ttl = it.get("t", "")
        num = it.get("n")
        if num is not None:
            bs = 0.235
            rect(slide, tx, ty + 0.03, bs, bs, acc, None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.18)
            s = slide.shapes[-1]
            stf = s.text_frame
            para(stf, num, size=10.5, color="FFFFFF", bold=True, align="c", first=True, font=MONO)
            stf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tx += bs + 0.13
            tw_ -= bs + 0.13
        if ttl:
            tb, tf = add_tb(slide, tx, ty, tw_, 0.3)
            para(tf, ttl, size=tsize, color=INK, bold=True, first=True, lh=1.15)
            sub = it.get("s")
            if sub:
                tb, tf = add_tb(slide, tx, ty + 0.30, tw_, 0.2)
                para(tf, sub, size=9.5, color=GREY_T, first=True, font=MONO)
                ty += 0.30
            ty += 0.34
        body = it.get("b")
        if body:
            avail = cy + ch - ty - 0.16
            if isinstance(body, (list, tuple)):
                yy = ty + 0.03
                step = (cy + ch - 0.14 - yy) / len(body)
                for ln in body:
                    pre, txt = ("", ln) if isinstance(ln, str) else ln
                    mark = {"g": "✓ ", "w": "! "}.get(pre, "· ")
                    lines, sz, _ = fit_wrap(mark + txt, tw_, step - 0.02, bsize, lh=1.30)
                    tb, tf = add_tb(slide, tx + 0.02, yy, tw_ - 0.02, step)
                    para(tf, chr(10).join(lines), size=sz, color=INK_SOFT, first=True, lh=1.30)
                    yy += step
            else:
                lines, sz, used = fit_wrap(body, tw_, avail, bsize, lh=1.34)
                tb, tf = add_tb(slide, tx, ty + 0.02, tw_, avail)
                for j, ln in enumerate(lines):
                    para(tf, ln, size=sz, color=INK_SOFT, first=(j == 0), lh=1.34)


def bullets(slide, items, bounds, size=13.0, lh=1.40, mark="—", mcol=None, gap=0.0):
    x, y, w, h = bounds
    if h < 0.34:
        h = 0.34
    n = len(items)
    step = (h - gap * (n - 1)) / n
    for i, it in enumerate(items):
        if isinstance(it, (list, tuple)):
            ttl, txt = it[0], it[1]
        else:
            ttl, txt = None, it
        yy = y + i * (step + gap)
        xx = x
        tb, tf = add_tb(slide, xx, yy, w, step)
        col = mcol or INK_SOFT
        if ttl:
            n1 = len(wrap(mark + " " + ttl, w - 0.24, size, True))
            h1 = n1 * size * 1.28 / 72.0 + 0.02
            lines, sz, _ = fit_wrap(mark + " " + ttl, w - 0.34, h1, size, lh=1.28, bold=True)
            para(tf, chr(10).join(lines), size=sz, color=INK, bold=True, first=True, lh=1.28)
            h2 = max(0.16, step - h1 - 0.02)
            lines2, sz2, _ = fit_wrap(txt, w - 0.24, h2, size - 1.4, lh=1.26)
            para(tf, chr(10).join(lines2), size=sz2, color=INK_SOFT, lh=1.30, before=2)
        else:
            lines, sz, used = fit_wrap(mark + " " + txt, w - 0.34, step + 0.02, size, lh=1.28)
            para(tf, chr(10).join(lines), size=sz, color=col, first=True, lh=1.28)


def flow(slide, steps, bounds, size=10.2, tsize=12.2, accent=INK_SOFT, per_row=4):
    x, y, w, h = bounds
    rows = -(-len(steps) // per_row)
    gapy = 0.22
    rh = (h - gapy * (rows - 1)) / rows
    for i, st in enumerate(steps):
        r_ = i // per_row
        idx = i % per_row
        if r_ % 2 == 1:
            idx = per_row - 1 - idx  # 蛇形走向
        cw = (w - 0.16 * (per_row - 1)) / per_row
        cx, cy = x + idx * (cw + 0.16), y + r_ * (rh + gapy)
        s = rect(slide, cx, cy, cw, rh, tint(accent, k=0.075), tint(accent, k=0.30), 0.75,
                 MSO_SHAPE.PENTAGON, adj=0.16)
        tf = s.text_frame
        tf.margin_left, tf.margin_right = Inches(0.16), Inches(0.20)
        tf.margin_top = Inches(0.12)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        num, ttl, body = st
        para(tf, "%s  %s" % (num, ttl), size=tsize, color=INK, bold=True, first=True, lh=1.15)
        lines, sz, used = fit_wrap(body, cw - 0.42, rh - 0.52, size)
        para(tf, "\n".join(lines), size=sz, color=INK_SOFT, lh=1.28, before=3)
        ay = cy + rh - 0.30
        if r_ % 2 == 0 and i % per_row < per_row - 1:
            ar = rect(slide, cx + cw - 0.09, ay, 0.18, 0.16, accent, None,
                      shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            ar.rotation = 90
        elif r_ % 2 == 1 and i % per_row > 0:
            ar = rect(slide, cx - 0.09, ay, 0.18, 0.16, accent, None,
                      shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            ar.rotation = 270


def table(slide, headers, rows, bounds, colw=None, hsize=11.5, bsize=11.0,
          accent=INK_SOFT, zebra=True, align=None):
    x, y, w, h = bounds
    n = len(headers)
    colw = colw or [w / n] * n
    tot = sum(colw)
    colw = [c * w / tot for c in colw]
    hh = 0.38
    rh = (h - hh) / max(len(rows), 1)
    rect(slide, x, y, w, hh, accent)
    cx = x
    for j, hd in enumerate(headers):
        tb, tf = add_tb(slide, cx + 0.10, y, colw[j] - 0.16, hh)
        para(tf, hd, size=hsize, color="FFFFFF", bold=True, first=True, lh=1.1,
             align=(align[j] if align else "l"))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cx += colw[j]
    for i, row in enumerate(rows):
        ry = y + hh + i * rh
        if zebra and i % 2 == 0:
            rect(slide, x, ry, w, rh, "FFFFFF", None)
        else:
            rect(slide, x, ry, w, rh, tint(accent, k=0.05), None)
        rect(slide, x, ry + rh, w, 0.008, "E7E4DC")
        cx = x
        for j, cell in enumerate(row):
            col = INK_SOFT
            txt = cell
            if isinstance(cell, (list, tuple)):
                txt, col = cell
            lines, sz, used = fit_wrap(txt, colw[j] - 0.20, rh - 0.10, bsize, lh=1.28,
                                       min_size=8.4)
            ty0 = ry + max(0.03, (rh - 0.10 - len(lines) * sz * 1.28 / 72.0) / 2 + 0.02)
            tb, tf = add_tb(slide, cx + 0.10, ty0, colw[j] - 0.16, rh - 0.10)
            para(tf, chr(10).join(lines), size=sz, color=col, first=True, lh=1.28,
                 align=(align[j] if align else "l"))
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            cx += colw[j]


def quote_block(slide, text, bounds, size=13.0, accent=INK_SOFT, label=None):
    x, y, w, h = bounds
    card(slide, x, y, w, h, tint(accent, k=0.07), None)
    rect(slide, x, y, 0.055, h, accent)
    lines, sz, used = fit_wrap(text, w - 0.55, h - 0.42 if label else h - 0.30, size)
    tb, tf = add_tb(slide, x + 0.34, y + (0.30 if label else 0.14), w - 0.62, h)
    if label:
        para(tf, label, size=9.5, color=accent, bold=True, first=True, spc=1.0)
    for i, ln in enumerate(lines):
        para(tf, ln, size=sz, color=INK_SOFT, first=(i == 0 and not label), lh=1.34)


def chip(slide, x, y, txt, color=INK_SOFT, size=9.5, fill=None, fg=None, pad=0.13):
    wd = tw(txt, size, True) + pad * 2
    rect(slide, x, y, wd, 0.26, fill or tint(color, k=0.14), None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.4)
    tb, tf = add_tb(slide, x + pad, y + 0.045, wd - 2 * pad + 0.2, 0.2)
    para(tf, txt, size=size, color=fg or color, bold=True, first=True)
    return x + wd


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])
