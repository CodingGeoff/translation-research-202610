# -*- coding: utf-8 -*-
import sys
from pptx import Presentation

p = r"d:\10_Workspace\translation-research-202610\翻译研究方法\第1组\问卷调查法正反例_学院模板版.pptx"
prs = Presentation(p)
print("slides:", len(prs.slides._sldIdLst))

# 检查每个 slide 的 layout 是否继承 logo + 格言（通过 layout 元素）
for i, slide in enumerate(prs.slides):
    ly = slide.slide_layout
    # 统计 layout 里的 pic 和文字
    pics = 0
    texts = []
    for sh in ly.shapes:
        if sh.shape_type == 13:  # PICTURE
            pics += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip()[:40])
    slide_texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            slide_texts.append(sh.text_frame.text.strip()[:30])
    print("  slide %02d | layout=%r | layout_pics=%d | layout_text=%s" % (i+1, ly.name, pics, texts))
    print("           | slide_shapes=%d | first=%r" % (len(slide.shapes), slide_texts[:2]))

# 完整提取文本到 UTF-8 文件
out = []
for i, slide in enumerate(prs.slides):
    out.append("=== SLIDE %d ===" % (i+1))
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
with open(r"d:\10_Workspace\translation-research-202610\翻译研究方法\第1组\_out_text.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("text dump written.")
