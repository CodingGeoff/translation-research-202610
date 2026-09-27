# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Emu

def hexcolor(c):
    try:
        if c is None:
            return None
        if c.type is None:
            return None
        return str(c.rgb)
    except Exception:
        return None

def dump(path, max_shapes=200):
    print("=" * 90)
    print("FILE:", path)
    prs = Presentation(path)
    print("slide size: %.2f x %.2f in" % (prs.slide_width / 914400, prs.slide_height / 914400))
    print("slides:", len(prs.slides._sldIdLst))
    # theme colors
    try:
        theme = prs.slide_masters[0].element.getroottree()
    except Exception:
        theme = None
    for idx, slide in enumerate(prs.slides):
        print("-" * 70)
        print("SLIDE", idx + 1, "| layout:", slide.slide_layout.name)
        # background
        try:
            bg = slide.background
            fill = bg.fill
            print("  bg type:", fill.type)
        except Exception as e:
            print("  bg err", e)
        n = 0
        for sh in slide.shapes:
            if n >= max_shapes:
                break
            n += 1
            kind = sh.shape_type
            txt = ""
            if sh.has_text_frame:
                txt = sh.text_frame.text.replace("\n", " / ")[:50]
            fillc = None
            try:
                fillc = hexcolor(sh.fill.fore_color)
            except Exception:
                pass
            pos = "x=%.2f y=%.2f w=%.2f h=%.2f" % (Emu(sh.left).inches if sh.left is not None else -1,
                                                     Emu(sh.top).inches if sh.top is not None else -1,
                                                     Emu(sh.width).inches if sh.width is not None else -1,
                                                     Emu(sh.height).inches if sh.height is not None else -1)
            print("  [%s] %s | %s | fill=%s | %r" % (kind, pos, sh.name[:30], fillc, txt))

if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
